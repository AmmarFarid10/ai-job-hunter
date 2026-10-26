# mypy: disable-error-code="no-untyped-call"
import html
import io
import unicodedata
from dataclasses import dataclass

import pymupdf
from docx import Document
from docx.shared import Inches, Pt

from app.domain.tailoring import ResumeDocument


@dataclass(frozen=True)
class RenderedResume:
    docx: bytes
    pdf: bytes
    pages: int


def paragraphs(document: ResumeDocument) -> list[tuple[str, str]]:
    rows = [("name", document.name), ("headline", document.headline)]
    contact = " | ".join(value for value in [document.location, document.email] if value)
    if contact:
        rows.append(("body", contact))
    rows.extend([("heading", "Professional Summary"), ("body", document.summary)])
    if document.skills:
        rows.extend(
            [
                ("heading", "Technical Skills"),
                ("body", ", ".join(row.name for row in document.skills)),
            ]
        )
    rows.append(("heading", "Professional Experience"))
    for row in document.experiences:
        rows.extend([("role", f"{row.title} | {row.employer}"), ("dates", row.dates)])
        rows.extend(("bullet", bullet.text) for bullet in row.bullets)
    if document.education:
        rows.append(("heading", "Education"))
        for education in document.education:
            rows.append(("role", f"{education.degree} | {education.institution}"))
            if education.dates:
                rows.append(("dates", education.dates))
    return rows


def normalized(text: str) -> str:
    return "".join(unicodedata.normalize("NFKC", text).replace("\u00ad", "").split())


def render_resume(document: ResumeDocument) -> RenderedResume:
    rows = paragraphs(document)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    doc.core_properties.author = document.name
    doc.core_properties.title = "Professional Resume"
    for kind, text in rows:
        if kind == "heading":
            paragraph = doc.add_heading(text, level=1)
        elif kind == "bullet":
            paragraph = doc.add_paragraph(text, style="List Bullet")
        else:
            paragraph = doc.add_paragraph(text)
            if kind in {"name", "headline", "role"}:
                paragraph.runs[0].bold = True
            if kind == "name":
                paragraph.runs[0].font.size = Pt(19)
            if kind in {"name", "headline", "role"}:
                paragraph.paragraph_format.keep_with_next = True
    docx_stream = io.BytesIO()
    doc.save(docx_stream)
    docx_bytes = docx_stream.getvalue()
    output = []
    for kind, text in rows:
        tag = "h1" if kind == "name" else "h2" if kind == "heading" else "p"
        prefix = "- " if kind == "bullet" else ""
        output.append(f'<{tag} class="{kind}">{html.escape(prefix + text)}</{tag}>')
    css = """
    body { font-family: sans-serif; font-size: 10.5pt; color: #111; }
    p { margin: 0 0 5pt; line-height: 1.15; }
    h1 { font-size: 19pt; margin: 0 0 5pt; }
    h2 { font-size: 12pt; margin: 10pt 0 5pt; }
    .role, .headline { font-weight: bold; page-break-after: avoid; }
    .dates { font-size: 10pt; page-break-after: avoid; }
    .bullet { margin-left: 10pt; }
    """
    # All source text is escaped; the only HTML/CSS comes from this template.
    story = pymupdf.Story(html="<html><body>" + "".join(output) + "</body></html>", user_css=css)
    pdf_stream = io.BytesIO()
    writer = pymupdf.DocumentWriter(pdf_stream)
    page = pymupdf.Rect(0, 0, 595, 842)
    area = pymupdf.Rect(42, 42, 553, 800)
    more, pages = True, 0
    try:
        while more:
            if pages >= 30:
                raise ValueError("Resume exceeds the 30-page rendering limit")
            device = writer.begin_page(page)
            more, _ = story.place(area)
            story.draw(device)
            writer.end_page()
            pages += 1
    finally:
        writer.close()
    pdf_bytes = pdf_stream.getvalue()
    # Reject clipped/missing text and font failures; both formats carry the same content.
    docx_text = normalized(
        " ".join(row.text for row in Document(io.BytesIO(docx_bytes)).paragraphs)
    )
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as rendered:
        pdf_text = normalized(" ".join(page.get_text(sort=True) for page in rendered))
    for _, text in rows:
        if normalized(text) not in docx_text or normalized(text) not in pdf_text:
            raise ValueError("Rendered resume text differs from the verified document")
    return RenderedResume(docx=docx_bytes, pdf=pdf_bytes, pages=pages)
